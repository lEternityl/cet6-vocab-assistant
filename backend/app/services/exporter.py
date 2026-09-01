from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import (
    QuestionSet,
    Sentence,
    SentenceTranslation,
    UserVocab,
    WordOccurrence,
    WordStat,
)
from .vocabulary import get_vocabulary


HEADER_FILL = PatternFill("solid", fgColor="17324D")
ACCENT_FILL = PatternFill("solid", fgColor="E9F3EE")


def _style_sheet(sheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def export_question_set(session: Session, question_set_id: int) -> tuple[BytesIO, str]:
    question_set = session.get(QuestionSet, question_set_id)
    if not question_set:
        raise ValueError("题目不存在")
    vocab = get_vocabulary()
    raw_stats = list(
        session.scalars(
            select(WordStat)
            .where(WordStat.question_set_id == question_set_id)
            .order_by(WordStat.study_frequency.desc(), WordStat.lemma)
        ).all()
    )
    grouped_stats: dict[str, dict] = {}
    for stat in raw_stats:
        word = grouped_stats.setdefault(
            stat.lemma,
            {
                "lemma": stat.lemma,
                "pos_values": [],
                "study_frequency": 0,
                "total_frequency": 0,
                "passage_frequency": 0,
            },
        )
        word["pos_values"].append(stat.pos)
        word["study_frequency"] += stat.study_frequency
        word["total_frequency"] += stat.total_frequency
        word["passage_frequency"] += stat.passage_frequency
    stats = sorted(
        grouped_stats.values(),
        key=lambda row: (-row["study_frequency"], row["lemma"]),
    )
    statuses_by_lemma: dict[str, list[str]] = {}
    for row in session.scalars(select(UserVocab)).all():
        statuses_by_lemma.setdefault(row.lemma, []).append(row.status)

    def word_status(lemma: str) -> str:
        values = set(statuses_by_lemma.get(lemma, []))
        if "mastered" in values:
            return "mastered"
        if "learning" in values:
            return "learning"
        if values == {"suspended"}:
            return "suspended"
        return "new"

    workbook = Workbook()
    summary = workbook.active
    summary.title = "词频统计"
    summary.append(
        [
            "单词",
            "音标",
            "词性",
            "释义",
            "学习范围词频",
            "全文词频",
            "阅读正文词频",
            "掌握状态",
            "例句",
            "例句翻译",
        ]
    )
    detail = workbook.create_sheet("例句明细")
    detail.append(["单词", "词性", "页码", "内容类型", "例句", "翻译"])

    for stat in stats:
        lemma = stat["lemma"]
        pos_label = " · ".join(sorted(set(stat["pos_values"])))
        entry = vocab.get(lemma)
        examples = session.execute(
            select(
                Sentence.id,
                Sentence.text,
                Sentence.content_type,
                func.min(WordOccurrence.id),
            )
            .join(WordOccurrence, WordOccurrence.sentence_id == Sentence.id)
            .where(
                WordOccurrence.question_set_id == question_set_id,
                WordOccurrence.lemma == lemma,
            )
            .group_by(Sentence.id)
            .order_by(Sentence.sentence_index)
        ).all()
        translations = {
            row.sentence_hash: row.translation
            for row in session.scalars(
                select(SentenceTranslation).where(
                    SentenceTranslation.sentence_hash.in_(
                        [
                            session.get(Sentence, example.id).normalized_hash
                            for example in examples
                        ]
                    )
                )
            ).all()
        }
        example_texts = [example.text for example in examples]
        example_translations = [
            translations.get(session.get(Sentence, example.id).normalized_hash, "")
            for example in examples
        ]
        summary.append(
            [
                lemma,
                entry.phonetic if entry else "",
                pos_label,
                entry.chinese if entry else "",
                stat["study_frequency"],
                stat["total_frequency"],
                stat["passage_frequency"],
                word_status(lemma),
                "\n".join(example_texts),
                "\n".join(example_translations),
            ]
        )
        for example, translation in zip(
            examples, example_translations, strict=False
        ):
            sentence = session.get(Sentence, example.id)
            page_number = sentence.block.page.page_index + 1
            detail.append(
                [
                    lemma,
                    pos_label,
                    page_number,
                    sentence.content_type,
                    sentence.text,
                    translation,
                ]
            )

    _style_sheet(summary, [18, 22, 12, 42, 16, 12, 16, 14, 72, 72])
    _style_sheet(detail, [18, 12, 10, 14, 90, 90])
    for row in summary.iter_rows(min_row=2):
        if row[7].value == "mastered":
            for cell in row:
                cell.fill = ACCENT_FILL

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    safe_title = question_set.title.replace("/", "-")
    return buffer, f"{safe_title}-词汇统计.xlsx"
